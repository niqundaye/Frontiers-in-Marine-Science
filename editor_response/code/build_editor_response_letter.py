from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "editor_response" / "01_Response_to_Editor_DRAFT.docx"
NAVY = RGBColor(23, 54, 93)
BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 31, 31)
MUTED = RGBColor(89, 89, 89)
WHITE = "FFFFFF"
LIGHT = "F2F4F7"
AMBER = "FFF2CC"


def set_font(run, name: str = "Calibri", size: float = 11, bold: bool | None = None, color: RGBColor = DARK, italic: bool = False) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure_table(table, widths: list[int]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_paragraph(doc, text: str, *, bold_start: str | None = None, italic: bool = False, after: float = 4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    if bold_start and text.startswith(bold_start):
        first = p.add_run(bold_start)
        set_font(first, bold=True)
        rest = p.add_run(text[len(bold_start):])
        set_font(rest, italic=italic)
    else:
        run = p.add_run(text)
        set_font(run, italic=italic)
    return p


def add_bullet(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    set_font(p.add_run(text))


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05
    for style_name, size, color, before, after in [
        ("Title", 23, NAVY, 0, 4),
        ("Subtitle", 13, MUTED, 0, 16),
        ("Heading 1", 16, BLUE, 14, 6),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, NAVY, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.add_run("SUPPLEMENTARY MATERIAL RESPONSE | DOI 10.3389/fmars.2026.1809036"), size=8.5, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.add_run("Response to Request for Original Input Data and Optimisation Logs")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Point-by-point clarification and auditable replacement materials")

    metadata = [
        ("To", "Editorial Office"),
        ("From", "Dr. Liu, on behalf of the authors"),
        ("Date", "11 September 2026"),
        ("Re", "Original 31-province inputs, source files, and 30-run optimisation records"),
        ("Status", "Draft for author confirmation before submission"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(f"{label}: "), bold=True)
        set_font(p.add_run(value))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    set_cell_text = p.add_run(
        "Essential disclosure: the requested historical matrix, source workbooks, and original run logs are not present in the retained materials currently accessible for this response. The replacement files supplied here are explicitly labelled and must not be described as recovered originals."
    )
    set_font(set_cell_text, bold=True, color=RGBColor(127, 96, 0))
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), AMBER)
    p_pr.append(shd)

    add_paragraph(doc, "Dear Editorial Office,")
    add_paragraph(doc, "Thank you for the opportunity to clarify the status of the supporting materials. We audited the manuscript, retained aggregate tables and figures, the current reproducibility repository, and the materials available to the team preparing this response.")
    add_paragraph(doc, "The original 31-province input/coefficient workbook and the original optimisation output directory were not found in those retained materials. The retained evidence does not allow us to state whether the absent files are stored elsewhere, were deleted, or were never exported as standalone files. We therefore cannot ethically represent a newly constructed matrix or newly generated run log as a historical original. The attached package provides a complete inventory of the missing items and the strongest auditable replacements that can be produced from the retained evidence.")

    doc.add_heading("1. Original 31-province coefficient/input matrix", level=1)
    add_paragraph(doc, "The contemporaneous matrix used to generate the article's reported numerical results is unavailable in the retained archive. In particular, no file contains the original province row order, the complete 31 x 4 x 2 allocation bounds, or the full province-specific production, processing, marketing, social, economic and ecological inputs.")
    add_paragraph(doc, "Replacement supplied: `03_Reconstructed_31_Province_Inputs.xlsx`, `06_PUBLIC_DATA_AND_REVERSE_CALIBRATION_METHOD.md`, and the CSV tables in `data/source/` and `data/reconstructed/`. The workbook now separates (i) direct 2024 observations transcribed for all 31 province-level regions from six official NBS yearbook tables and (ii) the processed/calibrated fields used to create all 248 region-sector-mode records. It includes the four production sectors, population, disposable income, wastewater pressure, freight, e-commerce, model coefficients, national constraints, formulas and QC checks.")
    add_paragraph(doc, "Direct observations are labelled `official public 2024 NBS transcription; not historical author input`; transformations and assumptions are labelled `calibrated reconstruction; not historical author input`. The province names follow the standard NBS ordering to make the replacement inspectable. The article does not disclose the original mapping between r01-r31 and province names, so that mapping remains an explicit reconstruction assumption.")

    doc.add_heading("2. Original data files or spreadsheets used to construct the matrix", level=1)
    add_paragraph(doc, "No contemporaneous raw workbook, cleaned workbook, coefficient-construction workbook, EWM calculation file, interpolation log, moving-average calculation file, or join/harmonisation file is present in the retained archive.")
    add_paragraph(doc, "Replacement supplied: `data/source/` contains exact transcriptions of article Tables 1-4, Ministry and World Bank/FAO extracts, and `nbs_2024_31_province_public_panel.csv`. The new panel transcribes six official China Statistical Yearbook 2025 tables for all 31 province-level regions: aquatic-product output, population, disposable income, wastewater pollutants, freight, and enterprise e-commerce. Its adjacent data dictionary retains every printed unit and direct official URL. These public files support independent plausibility checks and transparent rerunning but are not claimed to be the contemporaneous files used for the reported experiment.")
    add_paragraph(doc, "The former NBS A0407 portal endpoint returned HTTP 403 during a renewed automated check. We therefore used directly accessible official yearbook tables instead of third-party aggregators. Table 12-15 supplies the 2024 marine/freshwater and capture/aquaculture output required for a real 31-province production backbone. The temporal calibration to the paper's 2023 total, the fresh-sales/deep-processing split, and all external proxy transformations are documented formula by formula in `06_PUBLIC_DATA_AND_REVERSE_CALIBRATION_METHOD.md`.")

    doc.add_heading("3. Original 30-run optimisation outputs and performance metrics", level=1)
    add_paragraph(doc, "The original 30-run logs used to generate the reported figures and aggregate HV/IGD results are unavailable. No retained file contains the historical seed list, initial populations, generation-by-generation populations, objective and constraint arrays, final non-dominated fronts, HV normalisation/reference point, IGD reference front, or per-run metric values.")
    add_paragraph(doc, "Two clearly separated replacement layers are supplied:")
    add_bullet(doc, "`runs/article_figure3_calibrated/` contains 30 HV and 30 IGD values per algorithm generated deterministically from manuscript-reported distribution anchors. They reproduce the comparison plot structure only and are not recovered run measurements.")
    add_bullet(doc, "`runs/new_30run_surrogate/` contains 30 newly executed independent runs for each of IA-NSGA-III, NSGA-III, MOEA/D and NSGA-II. It includes the configuration snapshot, all generation logs, relocation events, final objectives and seven constraint residuals, all 248-dimensional decision vectors, run summaries, pooled reference front, metric definitions, and an independent HV/IGD recomputation check.")
    add_paragraph(doc, "The new surrogate experiment uses the official NBS 2024 province-sector pattern, reverse-calibrated to the paper's disclosed 2023 national total, plus explicitly labelled processing/marketing/social/economic/ecological proxies. It uses 48 individuals and 30 generations to provide a practical, fully executable audit path. It does not reproduce the article's 200-individual, 1,000-generation historical experiment, because the missing historical province inputs and metric-reference definitions are necessary for a strict rerun.")
    add_paragraph(doc, "The manuscript also contains an unresolved textual discrepancy: the abstract reports IGD = 0.012, whereas the Figure 2 discussion states a final IGD of 0.06. In the absence of the original logs and reference front, the retained evidence cannot determine which value reflects the historical calculation. The package preserves the 0.012 disclosed anchor used in the calibrated comparison data and does not silently choose between the conflicting statements.")

    doc.add_heading("4. Exact scope of unavailable files and variables", level=1)
    scope_intro = add_paragraph(doc, "The detailed row-by-row inventory is provided in `02_Material_Availability.csv`. The unavailable material comprises:")
    scope_intro.paragraph_format.keep_with_next = True
    for item in [
        "province production by year, four fishery sectors and two utilisation modes, including the original region order and allocation bounds;",
        "regional TAC values, regional cold-chain/processing capacity, cold-storage throughput, CCI moving averages, fleet-power limits, minimum-supply and flow-balance inputs;",
        "regional workforce, unit-income coefficients, marginal output values, processing/marketing values, supply-chain costs and ecological weights;",
        "digitalisation observations for VMS/AIS density, traceability adoption and temperature-monitoring coverage, together with normalisation bounds, entropy weights and final D_r values;",
        "raw/cleaned spreadsheets and transformation records for interpolation, moving-average completion, unit conversion, statistical calibration and harmonisation;",
        "historical experiment seeds, software/hardware environment, initial populations, generation populations, feasibility and constraint histories, reference-direction histories and final Pareto fronts;",
        "historical HV/IGD inputs, normalisation constants, HV reference point, IGD reference set, per-run values, sensitivity runs and ablation runs.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("Material map", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["Editorial request", "Historical status", "Material supplied"]
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        set_cell_fill(cell, "17365D")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(text), size=9.5, bold=True, color=RGBColor(255, 255, 255))
    rows = [
        ("31-province matrix", "Not found", "Official 31-province panel + reconstructed XLSX/CSV + derivation equations"),
        ("Source spreadsheets", "Not found", "Six NBS tables + article/Ministry/World Bank extracts + source catalogue"),
        ("Original 30-run logs", "Not found", "Calibrated plot series + 120 newly executed surrogate runs with full logs"),
        ("Exact unavailable list", "Documented", "02_Material_Availability.csv + author confirmation checklist"),
    ]
    for request, status, supplied in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, (request, status, supplied), strict=True):
            set_font(cell.paragraphs[0].add_run(text), size=9.5)
    configure_table(table, [2300, 1700, 5360])
    for row_index, row in enumerate(table.rows):
        if row_index and row_index % 2 == 0:
            for cell in row.cells:
                set_cell_fill(cell, LIGHT)

    doc.add_heading("Closing statement", level=1)
    add_paragraph(doc, "We recognise that reconstructed inputs and new surrogate runs do not replace the evidentiary value of the historical originals. They are supplied to make the present model structure, assumptions and calculations inspectable while accurately disclosing the limits of the retained record. If any historical files are subsequently recovered, we will provide them with versioned checksums and update the material inventory.")
    add_paragraph(doc, "Yours sincerely,")
    add_paragraph(doc, "Dr. Liu\nOn behalf of the authors", after=0)

    doc.add_heading("Sources cited in this response", level=1)
    sources = [
        "National Bureau of Statistics of China, Fisheries Statistical Survey System (2024): https://www.stats.gov.cn/fw/bmdcxmsp/bmzd/202407/t20240719_1955798.html",
        "Ministry of Agriculture and Rural Affairs, 2023 National Fishery Economic Statistics Communique: https://yyj.moa.gov.cn/gzdt/202407/t20240705_6458486.htm",
        "National Bureau of Statistics data portal, provincial aquatic-product indicator A0407: https://data.stats.gov.cn/easyquery.htm?cn=E0103&zb=A0407",
        "National Bureau of Statistics, China Statistical Yearbook 2025 navigation: https://www.stats.gov.cn/sj/ndsj/2025/left_.htm",
        "NBS Table 12-15, Output of Aquatic Products: https://www.stats.gov.cn/sj/ndsj/2025/html/E12-15.jpg",
        "NBS Tables 2-5, 6-18, 8-10, 16-13 and 16-39: https://www.stats.gov.cn/sj/ndsj/2025/html/E02-05.jpg ; https://www.stats.gov.cn/sj/ndsj/2025/html/E06-18.jpg ; https://www.stats.gov.cn/sj/ndsj/2025/html/E08-10.jpg ; https://www.stats.gov.cn/sj/ndsj/2025/html/E16-13.jpg ; https://www.stats.gov.cn/sj/ndsj/2025/html/E16-39.jpg",
        "Repository and code archive: https://github.com/niqundaye/Frontiers-in-Marine-Science",
    ]
    for source in sources:
        add_bullet(doc, source)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
